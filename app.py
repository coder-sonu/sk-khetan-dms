import os, csv, io, json, sqlite3, hashlib, secrets, datetime, re, uuid
from flask import Flask, request, redirect, session, jsonify, render_template, send_file, abort, send_from_directory

BASE=os.path.dirname(os.path.abspath(__file__))
DB=(os.path.join('/tmp','dms.sqlite3') if os.environ.get('VERCEL') else os.path.join(BASE,'data','dms.sqlite3'))
UPLOADS=(os.path.join('/tmp','dms_uploads') if os.environ.get('VERCEL') else os.path.join(BASE,'uploads'))
os.makedirs(UPLOADS,exist_ok=True)
os.makedirs(os.path.dirname(DB),exist_ok=True)
app=Flask(__name__)
app.secret_key=os.environ.get('DMS_SECRET_KEY','change-this-v8-secret')

SCHEMA='''
CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY AUTOINCREMENT,username TEXT UNIQUE,password_hash TEXT NOT NULL,role TEXT NOT NULL DEFAULT 'Viewer',active INTEGER NOT NULL DEFAULT 1,created_at TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS vehicles(id INTEGER PRIMARY KEY AUTOINCREMENT,vehicle_no TEXT UNIQUE NOT NULL,company TEXT,location TEXT,make_model TEXT,vehicle_type TEXT,driver TEXT,phone TEXT,chassis_no TEXT,engine_no TEXT,commission_date TEXT,current_status TEXT DEFAULT 'Running',active INTEGER DEFAULT 1,created_at TEXT NOT NULL,updated_at TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS daily_entries(id INTEGER PRIMARY KEY AUTOINCREMENT,entry_date TEXT NOT NULL,vehicle_id INTEGER NOT NULL,opening_km REAL DEFAULT 0,closing_km REAL DEFAULT 0,diesel_ltr REAL DEFAULT 0,diesel_amount REAL DEFAULT 0,status TEXT,remarks TEXT,created_by INTEGER,created_at TEXT NOT NULL,UNIQUE(entry_date,vehicle_id));
CREATE TABLE IF NOT EXISTS breakdowns(id INTEGER PRIMARY KEY AUTOINCREMENT,vehicle_id INTEGER NOT NULL,breakdown_date TEXT NOT NULL,category TEXT,problem TEXT,repair TEXT,cost REAL DEFAULT 0,downtime_hours REAL DEFAULT 0,status TEXT DEFAULT 'Open',closed_date TEXT,remarks TEXT,created_by INTEGER,created_at TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS services(id INTEGER PRIMARY KEY AUTOINCREMENT,vehicle_id INTEGER NOT NULL,service_date TEXT NOT NULL,service_type TEXT,service_km REAL DEFAULT 0,cost REAL DEFAULT 0,next_service_date TEXT,next_service_km REAL DEFAULT 0,status TEXT DEFAULT 'Completed',remarks TEXT,created_by INTEGER,created_at TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS audit_log(id INTEGER PRIMARY KEY AUTOINCREMENT,user_id INTEGER,action TEXT,entity TEXT,entity_id INTEGER,details TEXT,created_at TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY,value TEXT);
CREATE TABLE IF NOT EXISTS breakdown_photos(id INTEGER PRIMARY KEY AUTOINCREMENT,breakdown_id INTEGER NOT NULL,filename TEXT NOT NULL,original_name TEXT,uploaded_by INTEGER,created_at TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS notification_queue(id INTEGER PRIMARY KEY AUTOINCREMENT,kind TEXT NOT NULL,recipient TEXT,subject TEXT,body TEXT,status TEXT DEFAULT 'Pending',created_at TEXT NOT NULL,sent_at TEXT);
'''

def db():
    c=sqlite3.connect(DB); c.row_factory=sqlite3.Row; c.executescript(SCHEMA); return c

def now(): return datetime.datetime.now().isoformat(timespec='seconds')
def today(): return datetime.date.today().isoformat()
def hashpw(p,s=None):
    s=s or secrets.token_hex(16); return f'{s}${hashlib.sha256((s+p).encode()).hexdigest()}'
def checkpw(p,h):
    try:
        s,d=h.split('$',1); return secrets.compare_digest(hashpw(p,s).split('$',1)[1],d)
    except: return False

def audit(action,entity,eid=None,details=''):
    if 'uid' in session:
        c=db(); c.execute('INSERT INTO audit_log(user_id,action,entity,entity_id,details,created_at) VALUES(?,?,?,?,?,?)',(session['uid'],action,entity,eid,details,now())); c.commit(); c.close()

def seed_vehicle_master():
    seed=os.path.join(BASE,'vehicle_master_seed.xlsx')
    if not os.path.exists(seed): return 0
    try:
        from openpyxl import load_workbook
        wb=load_workbook(seed,read_only=True,data_only=True); c=db(); count=0
        for ws in wb.worksheets:
            vals=list(ws.values)
            if not vals: continue
            headers=[str(x or '').strip().lower() for x in vals[0]]; idx={h:i for i,h in enumerate(headers)}
            def g(row,*names):
                for n in names:
                    i=idx.get(n.lower())
                    if i is not None: return row[i]
                return ''
            for row in vals[1:]:
                vn=str(g(row,'Vehicle No.','Vehicle No','Vehicle Number')).strip()
                if not vn: continue
                t=now(); c.execute('''INSERT INTO vehicles(vehicle_no,company,location,make_model,vehicle_type,driver,phone,chassis_no,engine_no,current_status,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?) ON CONFLICT(vehicle_no) DO UPDATE SET company=excluded.company,location=excluded.location,make_model=excluded.make_model,vehicle_type=excluded.vehicle_type,chassis_no=excluded.chassis_no,engine_no=excluded.engine_no,updated_at=excluded.updated_at''',(vn,str(g(row,'Company Name','Company') or '').strip(),str(g(row,'Operational Location','Location') or '').strip(),str(g(row,'Make & Model','Make/Model','Model') or '').strip(),str(g(row,'Door No. / Unit ID','Vehicle Type','Type') or '').strip(),'','',str(g(row,'Chassis No.') or '').strip(),str(g(row,'Engine No.') or '').strip(),'Running',t,t)); count+=1
        c.commit(); c.close(); return count
    except Exception:
        return 0

def init():
    c=db();
    if not c.execute('SELECT 1 FROM users LIMIT 1').fetchone(): c.execute('INSERT INTO users(username,password_hash,role,created_at) VALUES(?,?,?,?)',('admin',hashpw('admin123'),'Admin',now()))
    c.commit(); c.close()
    if not one('SELECT 1 FROM vehicles LIMIT 1'): seed_vehicle_master()

def user():
    if 'uid' not in session:return None
    c=db(); r=c.execute('SELECT * FROM users WHERE id=? AND active=1',(session['uid'],)).fetchone(); c.close(); return r

def req_login():
    if not user(): abort(401)

def q(sql,args=()):
    c=db(); r=c.execute(sql,args).fetchall(); c.close(); return [dict(x) for x in r]

def one(sql,args=()):
    c=db(); r=c.execute(sql,args).fetchone(); c.close(); return dict(r) if r else None

# Vercel imports this module instead of running it as __main__.
# Initialize the ephemeral SQLite schema/default admin on every cold start.
# init() is idempotent because the schema uses IF NOT EXISTS and seed upserts.
init()

@app.route('/')
def root(): return redirect('/app' if user() else '/login')
@app.route('/login',methods=['GET','POST'])
def login():
    if request.method=='POST':
        u=request.form.get('username','').strip(); p=request.form.get('password','')
        r=one('SELECT * FROM users WHERE username=? AND active=1',(u,))
        if r and checkpw(p,r['password_hash']): session.update(uid=r['id'],username=r['username'],role=r['role']); audit('LOGIN','user',r['id']); return redirect('/app')
        return render_template('login.html',error='Invalid username or password')
    return render_template('login.html',error='')
@app.route('/logout')
def logout(): session.clear(); return redirect('/login')
@app.route('/app')
def ui(): req_login(); return render_template('app.html',user=session['username'],role=session['role'],session_uid=session['uid'])

@app.get('/api/me')
def me(): req_login(); return jsonify(dict(user()))
@app.get('/api/health')
def health(): req_login(); return jsonify({'ok':True,'db':os.path.exists(DB),'server_time':now(),'version':'V12 Sk Khetan Group Fleet Management','features':['Mobile Daily Entry','Breakdown Photos','Breakdown SLA','Service Alerts','Vehicle 360','Diesel Analytics','MIS Export','Audit Trail','Sk Khetan Group Copilot']})
@app.get('/api/lists')
def lists(): req_login(); return jsonify({'companies':[x['company'] for x in q("SELECT DISTINCT company FROM vehicles WHERE company IS NOT NULL AND company<>'' ORDER BY company")],'locations':[x['location'] for x in q("SELECT DISTINCT location FROM vehicles WHERE location IS NOT NULL AND location<>'' ORDER BY location")]})
@app.get('/api/vehicles')
def vehicles(): req_login(); return jsonify(q('SELECT * FROM vehicles WHERE active=1 ORDER BY vehicle_no'))
@app.post('/api/vehicles')
def add_vehicle():
    req_login(); d=request.json or {}; vn=d.get('vehicle_no','').strip()
    if not vn:return jsonify(ok=False,error='Vehicle No required'),400
    c=db()
    try:
        t=now(); c.execute('INSERT INTO vehicles(vehicle_no,company,location,make_model,vehicle_type,driver,phone,chassis_no,engine_no,commission_date,current_status,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)',(vn,d.get('company'),d.get('location'),d.get('make_model'),d.get('vehicle_type'),d.get('driver'),d.get('phone'),d.get('chassis_no'),d.get('engine_no'),d.get('commission_date'),d.get('current_status','Running'),t,t)); vid=c.execute('SELECT last_insert_rowid()').fetchone()[0]; c.commit(); c.close(); audit('CREATE','vehicle',vid,vn); return jsonify(ok=True,id=vid)
    except sqlite3.IntegrityError: c.close(); return jsonify(ok=False,error='Vehicle already exists'),409

@app.post('/api/import')
def import_master():
    req_login();
    if session['role'] not in ('Admin','Manager'): return jsonify(ok=False,error='Admin/Manager only'),403
    f=request.files.get('file');
    if not f:return jsonify(ok=False,error='File required'),400
    name=f.filename.lower(); rows=[]
    if name.endswith('.csv'): rows=list(csv.DictReader(io.StringIO(f.read().decode('utf-8-sig'))))
    elif name.endswith(('.xlsx','.xlsm')):
        try:
            from openpyxl import load_workbook
            wb=load_workbook(f,read_only=True,data_only=True); ws=wb.active; vals=list(ws.values); headers=[str(x or '').strip() for x in vals[0]]
            rows=[dict(zip(headers,r)) for r in vals[1:] if any(r)]
        except Exception as e:return jsonify(ok=False,error=str(e)),400
    else:return jsonify(ok=False,error='Use CSV/XLSX/XLSM'),400
    def get(r,*names):
        low={str(k).strip().lower():v for k,v in r.items()}
        for n in names:
            if n.lower() in low:return low[n.lower()]
        return ''
    c=db(); n=0
    for r in rows:
        vn=str(get(r,'Vehicle No','Vehicle No.','Vehicle Number','Registration No')).strip()
        if not vn: continue
        vals=(vn,get(r,'Company'),get(r,'Location'),get(r,'Make/Model','Model'),get(r,'Vehicle Type','Type'),get(r,'Driver'),get(r,'Driver Phone','Phone'),get(r,'Chassis No','Chassis'),get(r,'Engine No','Engine'))
        t=now(); c.execute('''INSERT INTO vehicles(vehicle_no,company,location,make_model,vehicle_type,driver,phone,chassis_no,engine_no,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?) ON CONFLICT(vehicle_no) DO UPDATE SET company=excluded.company,location=excluded.location,make_model=excluded.make_model,vehicle_type=excluded.vehicle_type,driver=excluded.driver,phone=excluded.phone,chassis_no=excluded.chassis_no,engine_no=excluded.engine_no,updated_at=excluded.updated_at''',vals+(t,t)); n+=1
    c.commit(); c.close(); audit('IMPORT','vehicle',None,f'{n} rows'); return jsonify(ok=True,count=n)

@app.post('/api/daily')
def daily_add():
    req_login(); d=request.json or {}; v=one('SELECT id FROM vehicles WHERE vehicle_no=?',(d.get('vehicle_no'),))
    if not v:return jsonify(ok=False,error='Vehicle not found'),404
    try:
        c=db(); c.execute('INSERT INTO daily_entries(entry_date,vehicle_id,opening_km,closing_km,diesel_ltr,diesel_amount,status,remarks,created_by,created_at) VALUES(?,?,?,?,?,?,?,?,?,?)',(d.get('entry_date') or today(),v['id'],float(d.get('opening_km') or 0),float(d.get('closing_km') or 0),float(d.get('diesel_ltr') or 0),float(d.get('diesel_amount') or 0),d.get('status','Running'),d.get('remarks',''),session['uid'],now())); c.execute('UPDATE vehicles SET current_status=?,updated_at=? WHERE id=?',(d.get('status','Running'),now(),v['id'])); c.commit(); eid=c.execute('SELECT last_insert_rowid()').fetchone()[0]; c.close(); audit('CREATE','daily',eid,d.get('vehicle_no')); return jsonify(ok=True,id=eid)
    except sqlite3.IntegrityError:return jsonify(ok=False,error='Daily entry already exists for this vehicle/date'),409
@app.get('/api/daily')
def daily_list():
    req_login(); return jsonify(q('''SELECT d.*,v.vehicle_no,v.company,v.location,u.username created_by_name FROM daily_entries d JOIN vehicles v ON v.id=d.vehicle_id LEFT JOIN users u ON u.id=d.created_by ORDER BY d.entry_date DESC,d.id DESC'''))
@app.get('/api/missing')
def missing():
    req_login(); dt=request.args.get('date',today()); return jsonify(q('''SELECT v.vehicle_no,v.company,v.location FROM vehicles v LEFT JOIN daily_entries d ON d.vehicle_id=v.id AND d.entry_date=? WHERE v.active=1 AND d.id IS NULL ORDER BY v.vehicle_no''',(dt,)))

@app.post('/api/breakdowns')
def bd_add():
    req_login(); d=request.json or {}; v=one('SELECT id FROM vehicles WHERE vehicle_no=?',(d.get('vehicle_no'),));
    if not v:return jsonify(ok=False,error='Vehicle not found'),404
    c=db(); c.execute('INSERT INTO breakdowns(vehicle_id,breakdown_date,category,problem,repair,cost,downtime_hours,status,remarks,created_by,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?)',(v['id'],d.get('breakdown_date') or today(),d.get('category'),d.get('problem'),d.get('repair'),float(d.get('cost') or 0),float(d.get('downtime') or 0),d.get('status','Open'),d.get('remarks'),session['uid'],now())); c.execute("UPDATE vehicles SET current_status='Breakdown',updated_at=? WHERE id=?",(now(),v['id'])); c.commit(); eid=c.execute('SELECT last_insert_rowid()').fetchone()[0]; c.close(); audit('CREATE','breakdown',eid,d.get('vehicle_no')); return jsonify(ok=True,id=eid)
@app.get('/api/breakdowns')
def bd_list(): req_login(); return jsonify(q('''SELECT b.*,v.vehicle_no,v.company,v.location FROM breakdowns b JOIN vehicles v ON v.id=b.vehicle_id ORDER BY b.breakdown_date DESC,b.id DESC'''))
@app.post('/api/breakdowns/<int:eid>/close')
def bd_close(eid):
    req_login(); c=db(); r=c.execute('SELECT vehicle_id FROM breakdowns WHERE id=?',(eid,)).fetchone();
    if not r:c.close();return jsonify(ok=False,error='Not found'),404
    c.execute("UPDATE breakdowns SET status='Closed',closed_date=?,remarks=COALESCE(remarks,'')||? WHERE id=?",(today(),'\nClosed by '+session['username'],eid)); c.execute("UPDATE vehicles SET current_status='Running',updated_at=? WHERE id=?",(now(),r['vehicle_id'])); c.commit();c.close();audit('CLOSE','breakdown',eid);return jsonify(ok=True)

@app.post('/api/breakdowns/<int:eid>/photo')
def bd_photo(eid):
    req_login(); b=one('SELECT id FROM breakdowns WHERE id=?',(eid,))
    if not b:return jsonify(ok=False,error='Breakdown not found'),404
    f=request.files.get('photo')
    if not f or not f.filename:return jsonify(ok=False,error='Photo required'),400
    ext=os.path.splitext(f.filename)[1].lower()
    if ext not in ('.jpg','.jpeg','.png','.webp'):return jsonify(ok=False,error='Only JPG, PNG or WEBP'),400
    name=uuid.uuid4().hex+ext; f.save(os.path.join(UPLOADS,name))
    c=db();c.execute('INSERT INTO breakdown_photos(breakdown_id,filename,original_name,uploaded_by,created_at) VALUES(?,?,?,?,?)',(eid,name,f.filename,session['uid'],now()));c.commit();pid=c.execute('SELECT last_insert_rowid()').fetchone()[0];c.close();audit('UPLOAD','breakdown_photo',pid,str(eid));return jsonify(ok=True,id=pid,url='/uploads/'+name)

@app.get('/api/breakdowns/<int:eid>/photos')
def bd_photos(eid):
    req_login();return jsonify(q('SELECT id,breakdown_id,filename,original_name,created_at FROM breakdown_photos WHERE breakdown_id=? ORDER BY id DESC',(eid,)))

@app.get('/uploads/<path:name>')
def uploads(name):
    req_login(); return send_from_directory(UPLOADS,name)

@app.post('/api/services')
def svc_add():
    req_login(); d=request.json or {}; v=one('SELECT id FROM vehicles WHERE vehicle_no=?',(d.get('vehicle_no'),));
    if not v:return jsonify(ok=False,error='Vehicle not found'),404
    c=db(); c.execute('INSERT INTO services(vehicle_id,service_date,service_type,service_km,cost,next_service_date,next_service_km,status,remarks,created_by,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?)',(v['id'],d.get('service_date') or today(),d.get('service_type'),float(d.get('service_km') or 0),float(d.get('cost') or 0),d.get('next_service_date'),float(d.get('next_service_km') or 0),d.get('status','Completed'),d.get('remarks'),session['uid'],now()));c.commit();eid=c.execute('SELECT last_insert_rowid()').fetchone()[0];c.close();audit('CREATE','service',eid,d.get('vehicle_no'));return jsonify(ok=True,id=eid)
@app.get('/api/services')
def svc_list(): req_login(); return jsonify(q('''SELECT s.*,v.vehicle_no,v.company,v.location FROM services s JOIN vehicles v ON v.id=s.vehicle_id ORDER BY s.next_service_date IS NULL,s.next_service_date ASC,s.service_date DESC'''))

@app.get('/api/dashboard')
def dashboard():
    req_login(); dt=request.args.get('date',today()); co=request.args.get('company',''); lo=request.args.get('location',''); args=[dt]
    filt='';
    if co:filt+=' AND v.company=?';args.append(co)
    if lo:filt+=' AND v.location=?';args.append(lo)
    total=one('SELECT COUNT(*) n FROM vehicles v WHERE active=1'+(' AND company=?' if co else '')+(' AND location=?' if lo else ''),([co] if co else [])+([lo] if lo else []))['n']
    statuses=q('SELECT current_status,COUNT(*) n FROM vehicles v WHERE active=1'+(' AND company=?' if co else '')+(' AND location=?' if lo else '')+' GROUP BY current_status',([co] if co else [])+([lo] if lo else [])); sm={x['current_status']:x['n'] for x in statuses}
    td=one('''SELECT COALESCE(SUM(d.closing_km-d.opening_km),0) km,COALESCE(SUM(d.diesel_ltr),0) fuel FROM daily_entries d JOIN vehicles v ON v.id=d.vehicle_id WHERE d.entry_date=?'''+filt,args)['km']
    row=one('''SELECT COALESCE(SUM(d.diesel_ltr),0) fuel,COALESCE(SUM(d.closing_km-d.opening_km),0) km FROM daily_entries d JOIN vehicles v ON v.id=d.vehicle_id WHERE d.entry_date BETWEEN date(?,'-29 day') AND ?'''+filt,[dt,dt]+(([co] if co else [])+([lo] if lo else [])))
    return jsonify(total_vehicles=total,running=sm.get('Running',0),idle=sm.get('Idle',0),breakdown=sm.get('Breakdown',0),today_km=float(td or 0),today_fuel=float(one('''SELECT COALESCE(SUM(d.diesel_ltr),0) fuel FROM daily_entries d JOIN vehicles v ON v.id=d.vehicle_id WHERE d.entry_date=?'''+filt,args)['fuel']),fuel30=float(row['fuel']),km30=float(row['km']),mileage30=(float(row['km'])/float(row['fuel']) if row['fuel'] else 0),missing=len(missing().get_json()),open_breakdowns=one("SELECT COUNT(*) n FROM breakdowns WHERE status='Open'")['n'])

@app.get('/api/companies')
def companies():
    req_login(); dt=request.args.get('date',today()); return jsonify(q('''SELECT v.company,COUNT(DISTINCT v.id) vehicles,COALESCE(SUM(d.closing_km-d.opening_km),0) km,COALESCE(SUM(d.diesel_ltr),0) diesel,CASE WHEN COALESCE(SUM(d.diesel_ltr),0)>0 THEN SUM(d.closing_km-d.opening_km)/SUM(d.diesel_ltr) ELSE 0 END mileage FROM vehicles v LEFT JOIN daily_entries d ON d.vehicle_id=v.id AND d.entry_date BETWEEN date(?,'-29 day') AND ? WHERE v.active=1 GROUP BY v.company ORDER BY diesel DESC''',(dt,dt)))
@app.get('/api/alerts')
def alerts():
    req_login();
    due=q("SELECT s.*,v.vehicle_no,v.company,v.location FROM services s JOIN vehicles v ON v.id=s.vehicle_id WHERE s.next_service_date IS NOT NULL AND s.next_service_date<=date(?,'+7 day') ORDER BY s.next_service_date",(today(),))
    overdue=[x for x in due if x['next_service_date']<today()]
    ranking=q("SELECT COALESCE(category,'Other') category,COUNT(*) cases,COALESCE(SUM(downtime_hours),0) hours,COALESCE(SUM(cost),0) cost FROM breakdowns GROUP BY COALESCE(category,'Other') ORDER BY cases DESC,hours DESC")
    dq=q("SELECT v.vehicle_no,v.company,'Negative KM' issue FROM daily_entries d JOIN vehicles v ON v.id=d.vehicle_id WHERE d.closing_km<d.opening_km UNION ALL SELECT v.vehicle_no,v.company,'Negative Diesel' issue FROM daily_entries d JOIN vehicles v ON v.id=d.vehicle_id WHERE d.diesel_ltr<0 LIMIT 100")
    return jsonify(open_breakdowns=q("SELECT b.*,v.vehicle_no,v.company FROM breakdowns b JOIN vehicles v ON v.id=b.vehicle_id WHERE b.status='Open' ORDER BY b.breakdown_date"),service_due=due,overdue_service=overdue,breakdown_ranking=ranking,data_quality=dq)

@app.get('/api/notifications')
def notifications():
    req_login(); return jsonify(q('SELECT * FROM notification_queue ORDER BY id DESC LIMIT 200'))

@app.post('/api/notifications/test')
def notification_test():
    req_login()
    if session['role']!='Admin': return jsonify(ok=False,error='Admin only'),403
    d=request.json or {}; c=db(); c.execute('INSERT INTO notification_queue(kind,recipient,subject,body,status,created_at) VALUES(?,?,?,?,?,?)',(d.get('kind','EMAIL'),d.get('recipient',''),d.get('subject','DMS Test'),d.get('body','DMS notification test'),'Queued',now())); c.commit(); eid=c.execute('SELECT last_insert_rowid()').fetchone()[0]; c.close(); audit('CREATE','notification',eid); return jsonify(ok=True,id=eid,status='Queued')

@app.get('/api/vehicle/<path:vn>')
def vehicle_360(vn):
    req_login(); v=one('SELECT * FROM vehicles WHERE vehicle_no=?',(vn,))
    if not v:return jsonify(master=None),404
    return jsonify(master=v,daily=q('SELECT * FROM daily_entries WHERE vehicle_id=? ORDER BY entry_date DESC,id DESC',(v['id'],)),breakdowns=q('SELECT * FROM breakdowns WHERE vehicle_id=? ORDER BY breakdown_date DESC,id DESC',(v['id'],)),services=q('SELECT * FROM services WHERE vehicle_id=? ORDER BY service_date DESC,id DESC',(v['id'],)))

@app.get('/api/users')
def users_list():
    req_login()
    if session['role']!='Admin': return jsonify(ok=False,error='Admin only'),403
    return jsonify(q('SELECT id,username,role,active,created_at FROM users ORDER BY username'))

@app.post('/api/users')
def user_create():
    req_login()
    if session['role']!='Admin': return jsonify(ok=False,error='Admin only'),403
    d=request.json or {}
    username=(d.get('username') or '').strip()
    password=d.get('password') or ''
    role=d.get('role') or 'Viewer'
    if not re.fullmatch(r'[A-Za-z0-9._-]{3,40}', username):
        return jsonify(ok=False,error='Login ID: 3-40 characters, letters/numbers/._- only'),400
    if len(password)<8: return jsonify(ok=False,error='Password must be at least 8 characters'),400
    if role not in ('Admin','Manager','Operator','Viewer'): return jsonify(ok=False,error='Invalid role'),400
    c=db()
    try:
        t=now(); c.execute('INSERT INTO users(username,password_hash,role,active,created_at) VALUES(?,?,?,?,?)',(username,hashpw(password),role,1,t)); uid=c.execute('SELECT last_insert_rowid()').fetchone()[0]; c.commit(); c.close(); audit('CREATE','user',uid,username); return jsonify(ok=True,id=uid,username=username,role=role)
    except sqlite3.IntegrityError:
        c.close(); return jsonify(ok=False,error='Login ID already exists'),409

@app.post('/api/users/<int:uid>/password')
def user_password(uid):
    req_login()
    if session['role']!='Admin' and session['uid']!=uid: return jsonify(ok=False,error='Not allowed'),403
    d=request.json or {}; password=d.get('password') or ''
    if len(password)<8: return jsonify(ok=False,error='Password must be at least 8 characters'),400
    c=db(); cur=c.execute('UPDATE users SET password_hash=? WHERE id=?',(hashpw(password),uid)); c.commit(); c.close()
    if cur.rowcount==0: return jsonify(ok=False,error='User not found'),404
    audit('PASSWORD_CHANGE','user',uid); return jsonify(ok=True)

@app.post('/api/users/<int:uid>/toggle')
def user_toggle(uid):
    req_login()
    if session['role']!='Admin': return jsonify(ok=False,error='Admin only'),403
    if uid==session['uid']: return jsonify(ok=False,error='You cannot deactivate your own login'),400
    r=one('SELECT active FROM users WHERE id=?',(uid,))
    if not r: return jsonify(ok=False,error='User not found'),404
    new=0 if r['active'] else 1; c=db(); c.execute('UPDATE users SET active=? WHERE id=?',(new,uid)); c.commit(); c.close(); audit('TOGGLE','user',uid,str(new)); return jsonify(ok=True,active=new)

@app.get('/api/assistant')
def assistant_get():
    req_login(); return jsonify(answer='POST a question to /api/assistant')
@app.post('/api/assistant')
def assistant():
    req_login(); text=(request.json or {}).get('text','').strip().lower()
    if not text:return jsonify(answer='Please enter a question.')
    if any(k in text for k in ['breakdown','break down']):
        data=q("SELECT b.breakdown_date,v.vehicle_no,v.company,v.location,b.problem,b.status FROM breakdowns b JOIN vehicles v ON v.id=b.vehicle_id WHERE b.status='Open' ORDER BY b.breakdown_date DESC")
        return jsonify(answer={'type':'breakdown','count':len(data),'items':data[:50]})
    if 'service' in text or 'due' in text:
        data=q("SELECT s.next_service_date,v.vehicle_no,v.company,v.location FROM services s JOIN vehicles v ON v.id=s.vehicle_id WHERE s.next_service_date IS NOT NULL AND s.next_service_date<=date(?,'+7 day') ORDER BY s.next_service_date",(today(),))
        return jsonify(answer={'type':'service_due_7_days','count':len(data),'items':data[:50]})
    if 'missing' in text or 'entry' in text:
        dt=today(); data=q("SELECT v.vehicle_no,v.company,v.location FROM vehicles v LEFT JOIN daily_entries d ON d.vehicle_id=v.id AND d.entry_date=? WHERE v.active=1 AND d.id IS NULL ORDER BY v.vehicle_no",(dt,))
        return jsonify(answer={'type':'missing_daily_entry','date':dt,'count':len(data),'items':data[:100]})
    if 'diesel' in text or 'fuel' in text or 'km/l' in text or 'mileage' in text:
        row=one("SELECT COALESCE(SUM(closing_km-opening_km),0) km,COALESCE(SUM(diesel_ltr),0) diesel FROM daily_entries WHERE entry_date BETWEEN date(?,'-29 day') AND ?",(today(),today()))
        return jsonify(answer={'type':'diesel_30_days','km':float(row['km']),'diesel_ltr':float(row['diesel']),'km_per_ltr':float(row['km'])/float(row['diesel']) if row['diesel'] else 0})
    return jsonify(answer='Try: breakdown, service due, missing daily entry, diesel 30 days.')

@app.get('/api/backup')
def backup():
    req_login()
    if session['role']!='Admin': return jsonify(ok=False,error='Admin only'),403
    c=db(); tables=['vehicles','daily_entries','breakdowns','services','users','audit_log','notification_queue']; out={}
    for t in tables: out[t]=[dict(x) for x in c.execute('SELECT * FROM '+t).fetchall()]
    c.close(); raw=json.dumps(out,indent=2,default=str).encode(); return send_file(io.BytesIO(raw),as_attachment=True,download_name='dms_backup_'+today()+'.json',mimetype='application/json')

@app.get('/api/export/<name>')
def export_csv(name):
    req_login(); allowed={'vehicles':'SELECT * FROM vehicles ORDER BY vehicle_no','daily':'SELECT d.*,v.vehicle_no,v.company,v.location FROM daily_entries d JOIN vehicles v ON v.id=d.vehicle_id ORDER BY d.entry_date DESC','breakdowns':'SELECT b.*,v.vehicle_no,v.company,v.location FROM breakdowns b JOIN vehicles v ON v.id=b.vehicle_id ORDER BY b.breakdown_date DESC','services':'SELECT s.*,v.vehicle_no,v.company,v.location FROM services s JOIN vehicles v ON v.id=s.vehicle_id ORDER BY s.service_date DESC','audit':'SELECT a.*,u.username FROM audit_log a LEFT JOIN users u ON u.id=a.user_id ORDER BY a.id DESC'}
    if name not in allowed: abort(404)
    rows=q(allowed[name]); buf=io.StringIO(); w=csv.DictWriter(buf,fieldnames=(list(rows[0].keys()) if rows else ['message'])); w.writeheader()
    if rows:
        for r in rows:w.writerow(r)
    else:w.writerow({'message':'No records'})
    return send_file(io.BytesIO(buf.getvalue().encode('utf-8-sig')),as_attachment=True,download_name=name+'_'+today()+'.csv',mimetype='text/csv')

@app.get('/api/management-mis')
def management_mis():
    req_login(); dt=request.args.get('date',today())
    return jsonify(date=dt,summary=dashboard().json,companies=companies().json,alerts=alerts().json)

if __name__=='__main__':
    init(); app.run(host='0.0.0.0',port=int(os.environ.get('PORT','8000')),debug=False)
